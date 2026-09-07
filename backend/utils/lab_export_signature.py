from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Callable, Dict, Optional

from docx import Document
from docx.shared import Inches
from flask import current_app

DEFAULT_MANAGER_SIGNATURE_FILENAME = "Eeid_Charles-ch.jpg"


def _signature_roots() -> list[Path]:
    candidates = []
    configured = str(os.getenv("LAB_SIGNATURE_DIR") or "").strip()
    if configured:
        candidates.append(Path(configured))
    candidates.extend(
        [
            Path(r"C:\inetpub\wwwroot\newweb2021\pic\itac"),
            Path(r"D:\wwwroot\newweb2021\pic\itac"),
        ]
    )
    return candidates


def _signature_filename(values: Dict) -> str:
    approval = values.get("approval") or {}
    reference = approval.get("manager_signature_file") or values.get("manager_signature_file")
    if isinstance(reference, dict):
        reference = reference.get("name") or reference.get("path") or reference.get("url")
    text = str(reference or "").strip().replace("\\", "/")
    text = text.split("?", 1)[0].split("#", 1)[0]
    return os.path.basename(text) or DEFAULT_MANAGER_SIGNATURE_FILENAME


def _find_signature(values: Dict) -> Optional[Path]:
    filename = _signature_filename(values)
    if not filename:
        return None
    return next((root / filename for root in _signature_roots() if (root / filename).is_file()), None)


def _load_department_manager_filename() -> str:
    conn = current_app.config["GET_DB_CONN"]()
    try:
        row = conn.cursor().execute(
            """
SELECT TOP 1 ISNULL(sign_e, sign_c) AS signature
  FROM dbo.staff
 WHERE ISNULL(dep_manager, 0) = 1
 ORDER BY id
"""
        ).fetchone()
        reference = str(row[0] or "").strip() if row else ""
        reference = reference.split("?", 1)[0].split("#", 1)[0].replace("\\", "/")
        return os.path.basename(reference)
    finally:
        conn.close()


def _load_form_values(form_id: str) -> Dict:
    conn = current_app.config["GET_DB_CONN"]()
    try:
        row = conn.cursor().execute(
            """
SELECT TOP 1 values_json
  FROM dbo.myInstance
 WHERE CONVERT(varchar(36), form_id) = ?
""",
            (form_id,),
        ).fetchone()
        if not row or not row[0]:
            return {}
        return json.loads(row[0])
    finally:
        conn.close()


def _load_workflow_state(form_id: str) -> Dict:
    conn = current_app.config["GET_DB_CONN"]()
    try:
        row = conn.cursor().execute(
            """
SELECT TOP 1
       ISNULL(workflow_status, '') AS workflow_status,
       ISNULL(workflow_temp_no, '') AS workflow_temp_no,
       ISNULL(workflow_formal_no, '') AS workflow_formal_no,
       ISNULL(workflow_report_no, '') AS workflow_report_no
  FROM dbo.lab_workflow_state
 WHERE CONVERT(varchar(36), form_id) = ?
""",
            (form_id,),
        ).fetchone()
        if not row:
            return {}
        return {
            "workflow_status": str(row[0] or "").strip(),
            "workflow_temp_no": str(row[1] or "").strip(),
            "workflow_formal_no": str(row[2] or "").strip(),
            "workflow_report_no": str(row[3] or "").strip(),
        }
    except Exception:
        return {}
    finally:
        conn.close()


def _iter_doc_paragraphs(container):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_doc_paragraphs(cell)


def _replace_text_in_docx(docx_path: Path, replacements: Dict[str, str]) -> None:
    if not replacements:
        return
    document = Document(docx_path)
    changed = False
    for paragraph in _iter_doc_paragraphs(document):
        for run in paragraph.runs:
            text = run.text
            new_text = text
            for old, new in replacements.items():
                if old and old in new_text:
                    new_text = new_text.replace(old, new)
            if new_text != text:
                run.text = new_text
                changed = True
    if changed:
        document.save(docx_path)


def _normalize_entrust_number(docx_path: Path, form_id: str, values: Dict) -> None:
    workflow = _load_workflow_state(form_id)
    formal_no = str(workflow.get("workflow_formal_no") or values.get("workflow_formal_no") or "").strip()
    if not formal_no:
        return

    temp_candidates = [
        str(workflow.get("workflow_temp_no") or "").strip(),
        str(values.get("workflow_temp_no") or "").strip(),
        str(values.get("lab_no") or "").strip(),
        str(values.get("entrust_no") or "").strip(),
    ]
    temp_no = next(
        (
            candidate
            for candidate in temp_candidates
            if candidate
            and candidate != formal_no
            and candidate.startswith("T")
        ),
        "",
    )
    if not temp_no:
        return

    _replace_text_in_docx(docx_path, {temp_no: formal_no})


def _requires_manager_signature(values: Dict) -> bool:
    approval = values.get("approval") or {}
    return (
        str(values.get("status_code") or values.get("status") or "").strip()
        in {"2", "3", "4", "RECEIVED_LOGO", "RECEIVED_NOLOGO", "REVIEW_DONE"}
        or bool(approval.get("manager_approval_checked"))
    )


def _signature_check(form_id: str, kind: str = "manager") -> Dict:
    values = _load_form_values(form_id)
    workflow = _load_workflow_state(form_id)
    if workflow.get("workflow_status") in {"RECEIVED_LOGO", "RECEIVED_NOLOGO", "REVIEW_DONE"}:
        values = {**values, "status_code": workflow["workflow_status"]}
    required = True if kind == "fixed" else _requires_manager_signature(values)
    filename = (
        DEFAULT_MANAGER_SIGNATURE_FILENAME
        if kind == "fixed" and required
        else _load_department_manager_filename()
        if required
        else ""
    )
    signature_values = {
        **values,
        "manager_signature_file": filename,
        "approval": {
            **(values.get("approval") or {}),
            "manager_signature_file": filename,
        },
    }
    path = _find_signature(signature_values) if filename else None
    return {
        "required": required,
        "filename": filename,
        "path": str(path) if path else "",
        "checked": [str(root / filename) for root in _signature_roots()] if filename else [],
    }


def validate_export_signature(form_id: str, kind: str = "manager") -> Dict:
    """Check the selected signature before an export builder creates a Word file."""
    result = _signature_check(form_id, kind)
    if result["required"] and not result["path"]:
        if not result["filename"]:
            raise FileNotFoundError(
                "找不到 staff.dep_manager=1 員工的 sign_e/sign_c 簽名檔"
            )
        checked = ", ".join(result["checked"])
        raise FileNotFoundError(
            f"主管簽名檔尚未設定或不存在：{result['filename']}；已檢查：{checked}"
        )
    return result


def _insert_signature(docx_path: Path, signature_path: Path) -> None:
    document = Document(docx_path)

    # LQP-14-01 final report signature line.
    for paragraph in document.paragraphs:
        if "Lab. Director" in paragraph.text and "Signature" in paragraph.text:
            label = re.sub(r"_+", "", paragraph.text).rstrip()
            paragraph.clear()
            paragraph.add_run(label + " ")
            paragraph.add_run().add_picture(str(signature_path), width=Inches(1.35))
            document.save(docx_path)
            return

    # Backward-compatible fallback for the entrust form template.
    target_cell = None
    seen = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                if "受理人員" in cell.text:
                    target_cell = cell
                    break
            if target_cell is not None:
                break
        if target_cell is not None:
            break

    if target_cell is None:
        raise RuntimeError("找不到 Word 模板中的「受理人員」簽名欄位")

    paragraph = target_cell.paragraphs[0]
    if paragraph.text and not paragraph.text.endswith(" "):
        paragraph.add_run(" ")
    paragraph.add_run().add_picture(str(signature_path), width=Inches(1.35))
    document.save(docx_path)


def install_lab_export_signature(module, builder_name: str = "_make_export_docx") -> None:
    original: Callable[[str], Dict] = getattr(module, builder_name)
    if getattr(original, "_manager_signature_wrapped", False):
        return

    def wrapped(form_id: str) -> Dict:
        kind = "fixed" if module.__name__.endswith("lab_final_report") or builder_name == "_build_final_report" else "manager"
        signature_check = validate_export_signature(form_id, kind)
        result = original(form_id)
        values = _load_form_values(form_id)
        if module.__name__.endswith("lab_forms") and builder_name == "_make_export_docx":
            out_path = Path(str(result.get("out_path") or ""))
            if out_path.is_file():
                _normalize_entrust_number(out_path, form_id, values)
        if not signature_check["required"]:
            return result

        signature_path = Path(signature_check["path"])
        if not signature_check["path"]:
            filename = signature_check["filename"] or "(未記錄主管簽名檔)"
            checked = ", ".join(str(root / filename) for root in _signature_roots())
            raise FileNotFoundError(f"找不到主管簽名檔，已檢查：{checked}")

        out_path = Path(str(result.get("out_path") or ""))
        if not out_path.is_file():
            raise FileNotFoundError(f"找不到匯出 Word：{out_path}")
        _insert_signature(out_path, signature_path)
        return result

    wrapped._manager_signature_wrapped = True
    setattr(module, builder_name, wrapped)
